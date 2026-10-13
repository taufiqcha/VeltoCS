import argparse
import io
import zipfile

import numpy as np
import openpyxl

from data import DATA, SCAN, VOLTAMETRI

FOLDER = {"LC": "Alat Sensor Cadmium LC", "EM": "Alat Sensor CD Groundturth Estat"}
LEVEL_PPM = [2, 4, 6, 8, 10, 20, 40, 60, 80, 100, 200, 400, 600, 800, 1000]
JUMLAH_TITIK, JUMLAH_SCAN = 181, 10
KISI = np.round(np.arange(-1.40, -0.50 + 0.01 / 2, 0.01), 4)
BUANG_SCAN = 2
MIN_TITIK = 15
AMBANG_KOSONG = 0.02


def scan_mentah(baris, s):
    pasangan = [(r[2 * s], r[2 * s + 1]) for r in baris if r[2 * s] is not None and r[2 * s + 1] is not None]
    if not pasangan:
        return None, None
    return np.array([p[0] for p in pasangan], dtype=float), np.array([p[1] for p in pasangan], dtype=float)


def samakan_panjang(v, i):
    asal, tuju = np.linspace(0.0, 1.0, num=len(i)), np.linspace(0.0, 1.0, num=JUMLAH_TITIK)
    return np.interp(tuju, asal, v), np.interp(tuju, asal, i)


def baca_arsip(zf, alat):
    V, I, y, ok = [], [], [], []
    for level in LEVEL_PPM:
        awalan = f"{FOLDER[alat]}/{level}/"
        for nama in sorted(n for n in zf.namelist() if n.startswith(awalan) and n.endswith(".xlsx")):
            buku = openpyxl.load_workbook(io.BytesIO(zf.read(nama)), data_only=True)
            baris = list(buku[buku.sheetnames[0]].iter_rows(values_only=True))[2:]
            vs, is_ = np.zeros((JUMLAH_SCAN, JUMLAH_TITIK)), np.zeros((JUMLAH_SCAN, JUMLAH_TITIK))
            sah = np.zeros(JUMLAH_SCAN, dtype=bool)
            for s in range(JUMLAH_SCAN):
                v, i = scan_mentah(baris, s)
                if v is None or len(v) < 20:
                    continue
                vs[s], is_[s] = samakan_panjang(v, i)
                sah[s] = True
            V.append(vs)
            I.append(is_)
            y.append(level)
            ok.append(sah)
    return np.stack(V), np.stack(I), np.array(y), np.stack(ok)


def satu_arah(v, i, naik):
    arah = np.sign(np.diff(v))
    pilih = arah > 0 if naik else arah < 0
    if pilih.sum() < MIN_TITIK:
        return np.full(KISI.size, np.nan)
    indeks = np.zeros(v.size, dtype=bool)
    indeks[:-1] |= pilih
    indeks[1:] |= pilih
    vv, ii = v[indeks], i[indeks]
    urut = np.argsort(vv)
    vv, ii = vv[urut], ii[urut]
    vv, awal = np.unique(np.round(vv, 6), return_inverse=True)
    ii = np.bincount(awal, weights=ii) / np.bincount(awal)
    return np.interp(KISI, vv, ii, left=np.nan, right=np.nan)


def satu_pengukuran(V, I, ok):
    baris = [np.concatenate([satu_arah(V[s], I[s], False), satu_arah(V[s], I[s], True)])
             for s in range(BUANG_SCAN, V.shape[0]) if ok[s]]
    if not baris:
        return np.full(2 * KISI.size, np.nan)
    with np.errstate(all="ignore"):
        return np.nanmedian(np.vstack(baris), axis=0)


def tambal(*matriks):
    layak = np.ones(matriks[0].shape[1], dtype=bool)
    for X in matriks:
        layak &= np.isnan(X).mean(axis=0) <= AMBANG_KOSONG
    keluar = []
    for X in matriks:
        Z = X[:, layak].copy()
        median = np.nanmedian(Z, axis=0)
        b, k = np.where(np.isnan(Z))
        Z[b, k] = median[k]
        keluar.append(Z)
    return keluar, layak


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("arsip", help="01_raw_voltammograms.zip from Zenodo")
    arsip = parser.parse_args().arsip
    DATA.mkdir(exist_ok=True)
    X, y = {}, {}
    with zipfile.ZipFile(arsip) as zf:
        for alat in ("LC", "EM"):
            V, I, y[alat], ok = baca_arsip(zf, alat)
            np.savez_compressed(SCAN / f"scans_{alat}.npz", V=V, I=I, y=y[alat], ok=ok)
            X[alat] = np.vstack([satu_pengukuran(V[n], I[n], ok[n]) for n in range(V.shape[0])])
            print(f"{alat}: {V.shape[0]} measurements, {ok.sum()}/{ok.size} scans read", flush=True)
    (X_LC, X_EM), layak = tambal(X["LC"], X["EM"])
    sumbu = np.concatenate([KISI, KISI])[layak]
    cabang = np.concatenate([np.zeros(KISI.size), np.ones(KISI.size)])[layak]
    np.savez_compressed(VOLTAMETRI, X_LC=X_LC, y_LC=y["LC"], X_EM=X_EM, y_EM=y["EM"], sumbu=sumbu, cabang=cabang)
    print(f"{layak.sum()} of {layak.size} variables kept, written to {VOLTAMETRI}")


if __name__ == "__main__":
    main()
